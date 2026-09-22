#!/usr/bin/env python3
"""parse_resume.py — one-shot, parallel "what the ATS parser sees" report for the ats-beater skill.

Runs every independent extraction concurrently (pdftotext default + layout, pdffonts, glyph sizes,
ink margins) and then the stream analyses (contact block, headings/zones, dates, placeholders,
hyphen fusions, column-reorder suspects, optional keyword counting), and prints ONE JSON report
plus a short human digest. One tool call replaces eight to ten.

Usage:
  python3 parse_resume.py resume.pdf                       # PDF
  python3 parse_resume.py resume.tex                       # compiles with pdflatex (scratch copy), scores the PDF
  python3 parse_resume.py resume.docx                      # python-docx body + table/header/footer/textbox audit
  python3 parse_resume.py resume.pdf --terms terms.json    # keyword counting with your term table
  python3 parse_resume.py resume.pdf --jd posting.md       # auto-extract candidate terms from a posting (prune them!)
  python3 parse_resume.py resume.pdf --out report.json --quiet

terms.json shape (write it after reading the posting; JD auto-extraction is only a starting point):
  [{"term": "Python", "priority": "required", "variants": ["python"]},
   {"term": "RAG", "priority": "required", "variants": ["rag", "retrieval[- ]augmented generation"]},
   {"term": "CS degree", "priority": "required", "variants": ["(b\\.?s\\.?|m\\.?s\\.?|bachelor|master)[^\\n]{0,20}computer science"], "home_zone": "education"},
   {"term": "REST APIs", "priority": "preferred", "variants": ["rest(ful)? apis?"], "synonyms": ["\\bapi\\b"]}]
  priority: "title" | "required" | "preferred"  (weights 2 / 2 / 1). variants are case-insensitive regexes; a plain
  space between words also matches a line wrap or hyphen; overlapping variants count a phrase once.
  The report's keyword block is report["keywords"] with "terms" (per-term rows incl. the variants used), "score_of_30",
  "missing", "synonym_only", "single_zone", "over_cap", "at_cap_3" (terms that must stay out of new text).

Dependencies: poppler-utils (pdftotext, pdffonts, pdftoppm). Optional: pdfplumber (glyph sizes), Pillow (ink
margins), pdflatex (for .tex), python-docx (for .docx). Missing optional pieces are reported, not fatal.
"""
import argparse, collections, json, os, re, shutil, subprocess, sys, tempfile, time
from concurrent.futures import ThreadPoolExecutor

# ----------------------------------------------------------------------------- helpers
def run(cmd, timeout=120):
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return p.returncode, p.stdout, p.stderr
    except FileNotFoundError:
        return 127, "", f"{cmd[0]}: not installed"
    except subprocess.TimeoutExpired:
        return 124, "", f"{cmd[0]}: timed out"

def have(tool):
    return shutil.which(tool) is not None

CANONICAL = {
    "summary":    r"^\s*(professional\s+summary|summary|profile|summary of qualifications)\s*:?\s*$",
    "skills":     r"^\s*(technical\s+skills|skills|core\s+competencies|technologies)\s*:?\s*$",
    "experience": r"^\s*((professional|work|relevant|industry)\s+experience|experience|employment)\s*:?\s*$",
    "projects":   r"^\s*(projects|selected\s+projects|personal\s+projects)\s*:?\s*$",
    "education":  r"^\s*education\s*:?\s*$",
    "other":      r"^\s*(certifications?|publications?|leadership(\s*&\s*activities)?|activities|honors(\s*&\s*awards)?|awards|interests|languages|volunteer(ing)?|references)\s*:?\s*$",
}
ZONE_OF = {"summary": "summary", "skills": "skills", "experience": "experience", "projects": "experience",
           "education": "education", "other": "other"}
KEY_ZONES = ("summary", "skills", "experience")

MONTHS_FULL = "January|February|March|April|May|June|July|August|September|October|November|December"
MONTHS_ABBR = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
DATE_PATTERNS = [
    ("Month YYYY", rf"\b(?:{MONTHS_FULL})\s+\d{{4}}\b"),
    ("Mon YYYY",   rf"\b(?:{MONTHS_ABBR})\.?\s+\d{{4}}\b"),
    ("Sept YYYY (non-standard)", r"\bSept\.?\s+\d{4}\b"),
    ("MM/YYYY",    r"\b(?:0?[1-9]|1[0-2])/\d{4}\b"),
    ("MM/YY (ambiguous)", r"\b(?:0?[1-9]|1[0-2])/\d{2}\b(?!\d)"),
    ("Season YYYY", r"\b(?:Spring|Summer|Fall|Autumn|Winter)\s+\d{4}\b"),
    ("full calendar date", r"\b\d{1,2}/\d{1,2}/\d{2,4}\b"),
]
PLACEHOLDER_RE = re.compile(r"\[(?:N|X|TBD|[^\]]{0,40}\bfill\b[^\]]{0,40}|\?+)\]|\\fillin\{|\bTBD\b|\bXX+\b|\bLorem ipsum\b|\[[A-Za-z ]{3,40}\]", re.I)

# ----------------------------------------------------------------------------- source → PDF
def compile_tex(tex_path, workdir):
    if not have("pdflatex"):
        return None, "pdflatex not installed — compile the .tex yourself and pass the PDF"
    shutil.copy(tex_path, workdir)
    base = os.path.basename(tex_path)
    log = ""
    for _ in range(2):
        p = subprocess.run(["pdflatex", "-interaction=nonstopmode", base], cwd=workdir, capture_output=True, text=True, timeout=180)
        log = p.stdout
    pdf = os.path.join(workdir, os.path.splitext(base)[0] + ".pdf")
    if not os.path.exists(pdf):
        tail = "\n".join(l for l in log.splitlines() if l.startswith("!"))[:800]
        return None, f"pdflatex produced no PDF:\n{tail}"
    return pdf, None

def docx_report(path):
    try:
        import docx  # python-docx
    except ImportError:
        return None, "python-docx not installed (pip install python-docx)"
    d = docx.Document(path)
    body = "\n".join(p.text for p in d.paragraphs)
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    headers, footers = [], []
    for s in d.sections:
        headers += [p.text for p in s.header.paragraphs if p.text.strip()]
        footers += [p.text for p in s.footer.paragraphs if p.text.strip()]
    xml = d.element.xml
    textboxes = len(re.findall(r"<w:txbxContent", xml))
    images = len(re.findall(r"<pic:pic\b|<w:drawing\b", xml))
    # simulate the linear stream: body paragraphs, then table cells row-major
    stream = body + ("\n" + "\n".join("\t".join(r) for t in tables for r in t) if tables else "")
    return {"stream": stream, "layout": stream, "tables": tables, "header_text": headers,
            "footer_text": footers, "textboxes": textboxes, "images": images}, None

# ----------------------------------------------------------------------------- PDF extractions (run in parallel)
def x_pdftotext(pdf):
    rc, out, err = run(["pdftotext", "-enc", "UTF-8", pdf, "-"])
    return {"stream": out if rc == 0 else "", "error": None if rc == 0 else err.strip()}

def x_layout(pdf):
    rc, out, err = run(["pdftotext", "-enc", "UTF-8", "-layout", pdf, "-"])
    return {"layout": out if rc == 0 else "", "error": None if rc == 0 else err.strip()}

def x_fonts(pdf):
    rc, out, err = run(["pdffonts", pdf])
    if rc != 0:
        return {"fonts": [], "error": err.strip()}
    fonts = []
    for line in out.splitlines()[2:]:
        parts = line.split()
        if len(parts) < 6:
            continue
        # columns: name  type(1-3 words)  encoding  emb  sub  uni  object-id  gen
        name = parts[0]
        emb, sub, uni = parts[-5], parts[-4], parts[-3]
        ftype = " ".join(parts[1:-6])
        fonts.append({"name": name, "type": ftype, "embedded": emb == "yes", "unicode_map": uni == "yes"})
    return {"fonts": fonts, "error": None}

def x_pageinfo(pdf):
    rc, out, err = run(["pdfinfo", pdf])
    pages = None; size = None
    for line in out.splitlines():
        if line.startswith("Pages:"): pages = int(line.split()[1])
        if line.startswith("Page size:"): size = line.split(":", 1)[1].strip()
    return {"pages": pages, "page_size": size, "error": None if rc == 0 else err.strip()}

def x_glyphs(pdf):
    try:
        import pdfplumber
    except ImportError:
        return {"error": "pdfplumber not installed (pip install pdfplumber) — font sizes not measured"}
    sizes = collections.Counter(); families = set(); body = 0; total = 0
    with pdfplumber.open(pdf) as doc:
        for page in doc.pages:
            for c in page.chars:
                if not c["text"].strip():
                    continue
                s = round(c["size"], 1); sizes[s] += 1; total += 1
                fam = re.sub(r"^[A-Z]{6}\+", "", c["fontname"])
                fam = re.sub(r"[-,](Bold|Italic|BoldItalic|Regular|Oblique|MT|PS).*$", "", fam)
                fam = re.sub(r"(BX|TI|SY|MI|R)?\d+$", "", fam)  # CMBX12 → CMBX → CM
                families.add(fam)
    if not total:
        return {"error": "no glyphs found — image-only PDF?"}
    # body text = the most common sizes that together cover 60% of glyphs
    hist = sorted(sizes.items(), key=lambda kv: -kv[1])
    covered = 0; body_sizes = []
    for s, n in hist:
        body_sizes.append(s); covered += n
        if covered / total >= 0.6:
            break
    return {"min_body_pt": min(body_sizes), "size_histogram": {str(k): v for k, v in sorted(sizes.items())},
            "font_families": sorted(families), "error": None}

def x_ink_margins(pdf, workdir):
    if not have("pdftoppm"):
        return {"error": "pdftoppm not installed — margins not measured"}
    try:
        from PIL import Image
        import numpy as np
    except ImportError:
        return {"error": "Pillow/numpy not installed — margins not measured"}
    prefix = os.path.join(workdir, "ink")
    rc, out, err = run(["pdftoppm", "-r", "300", "-f", "1", "-l", "1", "-png", pdf, prefix])
    if rc != 0:
        return {"error": err.strip()}
    png = next((os.path.join(workdir, f) for f in os.listdir(workdir) if f.startswith("ink") and f.endswith(".png")), None)
    if not png:
        return {"error": "render failed"}
    a = np.array(Image.open(png).convert("L"))
    rows = np.where(a.min(axis=1) < 128)[0]; cols = np.where(a.min(axis=0) < 128)[0]
    if not len(rows):
        return {"error": "blank page"}
    return {"top_in": round(rows[0] / 300, 2), "bottom_in": round((a.shape[0] - rows[-1]) / 300, 2),
            "left_in": round(cols[0] / 300, 2), "right_in": round((a.shape[1] - cols[-1]) / 300, 2), "error": None}

# ----------------------------------------------------------------------------- stream analyses
def analyze_contact(stream):
    lines = [l for l in stream.splitlines()]
    nonblank = [l for l in lines if l.strip()]
    head = "\n".join(nonblank[:5])
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", head)
    phone = re.search(r"(\+?\d[\d\s().-]{8,}\d)", head)
    urls = re.findall(r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com|[\w-]+\.(?:dev|io|app|me|com))/?[\w./-]*", head)
    return {"name_line": nonblank[0] if nonblank else "", "email_in_first_5": bool(email), "phone_in_first_5": bool(phone),
            "urls_in_first_5": urls, "first_5_lines": nonblank[:5]}

def analyze_headings(stream):
    lines = stream.splitlines()
    headings, zones, zone = [], {z: [] for z in ("summary", "skills", "experience", "education", "other")}, "other"
    for i, line in enumerate(lines):
        hit = next((k for k, pat in CANONICAL.items() if re.match(pat, line, re.I)), None)
        if hit:
            zone = ZONE_OF[hit]
            headings.append({"line": i + 1, "text": line.strip(), "zone": zone})
            continue
        # a short Title-Case line alone that is NOT canonical → suspected creative heading
        if 1 <= len(line.split()) <= 4 and line.strip() == line.strip().title() and not re.search(r"[|•,:–—\d@/]", line) \
           and i > 0 and lines[i - 1].strip() == "" and i + 1 < len(lines) and lines[i + 1].strip() == "":
            nxt = next((x for x in lines[i + 2:i + 6] if x.strip()), "")
            entry_like = re.search(r",\s*[A-Z]{2}\b|\bRemote\b|\d{4}", nxt)  # "City, ST" or a date → an entry header, not a section
            if not entry_like:
                headings.append({"line": i + 1, "text": line.strip(), "zone": None, "suspect_non_canonical": True})
        zones[zone].append(line)
    found = {h["zone"] for h in headings if h["zone"]}
    return {"headings": headings, "has_summary": "summary" in found, "has_skills": "skills" in found,
            "has_experience": "experience" in found, "has_education": "education" in found,
            "non_canonical_suspects": [h["text"] for h in headings if h.get("suspect_non_canonical")]}, zones

def analyze_dates(stream):
    styles = {}
    for label, pat in DATE_PATTERNS:
        ms = re.findall(pat, stream)
        if ms:
            styles[label] = len(ms)
    full_no_may = len(re.findall(rf"\b(?:{MONTHS_FULL.replace('May|', '')})\s+\d{{4}}\b", stream))
    abbr_no_may = len(re.findall(rf"\b(?:{MONTHS_ABBR.replace('May|', '')})\.?\s+\d{{4}}\b", stream))
    numeric = styles.get("MM/YYYY", 0) + styles.get("MM/YY (ambiguous)", 0)
    styles_in_use = sum(1 for n in (full_no_may, abbr_no_may, numeric) if n) + (1 if styles.get("Sept YYYY (non-standard)") else 0)
    consistent = styles_in_use <= 1
    return {"styles": styles, "consistent_month_style": consistent, "present_used": bool(re.search(r"\bPresent\b", stream)),
            "flags": [k for k in styles if "ambiguous" in k or "non-standard" in k or "calendar" in k]}

def analyze_placeholders(stream):
    return [l.strip() for l in stream.splitlines() if PLACEHOLDER_RE.search(l)]

def analyze_hyphen_joins(stream, layout):
    """Every line-end hyphen in the layout text, and how the joined word landed in the parser stream.
    pdftotext de-hyphenates line breaks: right for a soft break (docu-/ments), wrong for a real compound
    (metadata-/regex → 'metadataregex'). The model judges which is which; both are listed."""
    lines = layout.splitlines()
    joins = []
    for i, l in enumerate(lines[:-1]):
        m = re.search(r"(\w+)-\s*$", l)
        if not m:
            continue
        m2 = re.match(r"\s*(\w+)", lines[i + 1])
        if not m2:
            continue
        w1, w2 = m.group(1), m2.group(1)
        joined, hy = w1 + w2, w1 + "-" + w2
        if re.search(r"\b" + re.escape(hy) + r"\b", stream, re.I):
            seen = "hyphenated"
        elif re.search(r"\b" + re.escape(joined) + r"\b", stream, re.I):
            seen = "joined"
        else:
            seen = "neither"
        joins.append({"hyphenated": hy, "joined": joined, "in_stream_as": seen,
                      "likely_compound": bool(re.search(r"[A-Z]", w2)) or len(w1) >= 4 and len(w2) >= 4 and seen == "joined"})
    return joins

def analyze_column_reorder(stream, layout):
    """Right-aligned fragments (layout lines with a wide gap) that land far from their left partner in the stream."""
    s_lines = [l.strip() for l in stream.splitlines()]
    suspects = []
    cursor = 0
    for l in layout.splitlines():
        m = re.match(r"^(\S.*?\S)\s{4,}(\S.*\S)\s*$", l)
        if not m:
            continue
        left, right = m.group(1).strip(), m.group(2).strip()
        li = next((i for i in range(cursor, len(s_lines)) if s_lines[i].startswith(left[:25])), None)
        if li is None:
            continue
        cursor = li
        ri = next((i for i in range(li + 1, len(s_lines)) if s_lines[i] == right or s_lines[i].startswith(right[:25])), None)
        if ri is None:
            continue
        gap = ri - li
        between = s_lines[li + 1:ri]
        crossed = any(re.match(p, b, re.I) for b in between for p in CANONICAL.values())
        if gap > 3 or crossed:
            suspects.append({"left": left, "right": right, "stream_gap_lines": gap, "crosses_heading": crossed})
    return suspects

# ----------------------------------------------------------------------------- keywords
TECH_VOCAB = [  # regexes, case-insensitive; multi-word entries allowed. Extend freely.
    r"python", r"java(?!script)", r"c\+\+", r"c#", r"javascript", r"typescript", r"golang|\bgo\b", r"rust", r"kotlin", r"swift",
    r"scala", r"ruby", r"php", r"\bsql\b", r"nosql", r"html", r"css", r"react", r"angular", r"vue", r"next\.?js", r"node\.?js",
    r"express", r"django", r"flask", r"fastapi", r"spring boot|spring", r"rails", r"\.net", r"graphql", r"rest(?:ful)? apis?",
    r"postgres(?:ql)?", r"mysql", r"mongodb", r"redis", r"neo4j", r"cypher", r"dynamodb", r"sqlite", r"elasticsearch", r"kafka",
    r"rabbitmq", r"\baws\b", r"azure", r"\bgcp\b|google cloud", r"docker", r"kubernetes", r"terraform", r"ansible", r"jenkins",
    r"github actions", r"gitlab", r"ci/cd", r"\bgit\b", r"linux", r"bash", r"machine learning", r"deep learning", r"\bnlp\b",
    r"\bllms?\b", r"large language models?", r"\brag\b", r"retrieval[- ]augmented generation", r"embeddings?",
    r"vector (?:search|database)s?", r"pgvector", r"langchain", r"llamaindex", r"openai", r"anthropic",
    r"fine[- ]tuning", r"prompt[- ]engineer(?:ing|ed)", r"pytorch", r"tensorflow", r"scikit-learn", r"pandas", r"numpy", r"spark", r"unit tests?", r"integration tests?",
    r"test automation", r"pytest", r"jest", r"agile", r"scrum", r"microservices", r"distributed systems", r"data structures",
    r"algorithms", r"object[- ]oriented", r"\boop\b", r"system design", r"computer science", r"bachelor'?s?|\bb\.?s\.?(?=\s|,|$)|\bb\.?sc\b", r"master'?s?|\bm\.?s\.?(?=\s|,|$)|\bm\.?sc\b",
    r"structured outputs?", r"evaluation pipelines?", r"grounding", r"code[- ]reviews?",
]

ALIASES = {  # acronym ↔ spelled-out pairs count as one term
    r"\bllms?\b": "LLM", r"large language models?": "LLM",
    r"\brag\b": "RAG", r"retrieval[- ]augmented generation": "RAG",
    r"ci/cd": "CI/CD", r"object[- ]oriented": "OOP", r"\boop\b": "OOP",
    r"vector (?:search|database)s?": "vector search / database", r"pgvector": "pgvector",
    r"rest(?:ful)? apis?": "REST APIs", r"bachelor'?s?|\bb\.?s\.?(?=\s|,|$)|\bb\.?sc\b": "Bachelor's degree", r"master'?s?|\bm\.?s\.?(?=\s|,|$)|\bm\.?sc\b": "Master's degree",
}
HOME_ZONE_TERMS = {"Bachelor's degree": "education", "Master's degree": "education", "computer science": "education"}

def extract_jd_terms(jd_text):
    """Candidate terms from a posting: responsibilities and requirements → required; nice-to-have → preferred.
    A term seen under both keeps the higher priority. Output is a starting point — prune and extend it."""
    lines = jd_text.splitlines()
    prio = "required"; terms = collections.OrderedDict()
    def bounded(pat):
        pat = re.sub(r"(?<=\w) (?=\w)", r"[\\s-]+", pat)
        return (r"(?<![\w-])" if re.match(r"[\w\\]", pat) else "") + "(?:" + pat + ")" + (r"(?![\w-])" if re.search(r"[\w)?]$", pat) else "")
    for line in lines:
        low = line.lower()
        if re.search(r"nice[- ]to[- ]have|preferred|bonus|a plus\b|good to have|optional", low) and len(low) < 60:
            prio = "preferred"
        elif re.search(r"requirement|qualification|must[- ]have|what you.?ll (need|do)|responsibilit|you have|you will|about the role", low) and len(low) < 60:
            prio = "required"
        for pat in TECH_VOCAB:
            for m in re.finditer(bounded(pat), line, re.I):
                key = ALIASES.get(pat, m.group(0).lower())
                if key in terms:
                    if prio == "required": terms[key]["priority"] = "required"
                    if bounded(pat) not in terms[key]["variants"]: terms[key]["variants"].append(bounded(pat))
                else:
                    t = {"term": key, "priority": prio, "variants": [bounded(pat)], "auto": True}
                    if key in HOME_ZONE_TERMS: t["home_zone"] = HOME_ZONE_TERMS[key]
                    terms[key] = t
    title = next((l.strip().lstrip("# ").strip() for l in lines if l.strip()), "")
    if title:
        core = re.split(r"[,(\u2014\-|]", title)[0].strip()
        if core:
            terms["__title__"] = {"term": f"{core} (job title)", "priority": "title", "variants": [bounded(re.escape(core).replace(r"\ ", r"\s+"))], "auto": True}
    return list(terms.values())

def _norm_pat(pat):
    """A plain space between words in a variant also matches a line wrap or a hyphen ("code review" ~ "code-review")."""
    return re.sub(r"(?<=\w) (?=\w)", r"[\\s-]+", pat)

def _unique_hits(patterns, text):
    """Number of non-overlapping matches across all variants (overlapping variants count a phrase once)."""
    spans = []
    for p in patterns:
        spans += [m.span() for m in re.finditer(p, text, re.I)]
    spans.sort()
    merged = []
    for a, b in spans:
        if merged and a < merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], b))
        else:
            merged.append((a, b))
    return merged

def count_terms(terms, zones):
    # line breaks become spaces so a phrase wrapped across lines still matches
    zone_text = {z: re.sub(r"\s*\n\s*", " ", "\n".join(ls)).strip() for z, ls in zones.items()}
    alltext = "\n".join(zone_text.values())
    PRESENCE = {0: 0.0, 1: 0.75}; ZONE_MULT = {0: 0.0, 1: 0.70, 2: 0.85, 3: 1.0}
    rows = []; tot_w = tot_credit = 0.0; stuffed = []
    for t in terms:
        w = 2 if t.get("priority") in ("title", "required") else 1
        pats = [_norm_pat(p) for p in (t.get("variants") or [re.escape(t["term"])])]
        def cnt(s):
            hits = _unique_hits(pats, s)
            n = len(hits)
            # 'Spelled-out (ACRONYM)' pairs count once
            for full, acro in re.findall(r"([A-Za-z][A-Za-z\- ]{6,60}) \(([A-Z]{2,6}s?)\)", s):
                if any(re.search(p, full, re.I) for p in pats) and any(re.search(p, acro, re.I) for p in pats):
                    n -= 1
            return max(n, 0)
        per = {z: cnt(zone_text.get(z, "")) for z in ("summary", "skills", "experience", "education", "other")}
        total = sum(per.values())
        syn = len(_unique_hits([_norm_pat(p) for p in t.get("synonyms", [])], alltext)) if t.get("synonyms") else 0
        home = t.get("home_zone")
        if home:
            credit = 1.0 if per.get(home, 0) else 0.0; zones_hit = None
        else:
            zones_hit = sum(1 for z in KEY_ZONES if per.get(z, 0))
            presence = PRESENCE.get(total, 1.0)
            credit = presence * ZONE_MULT[zones_hit]
            if total == 0 and syn:
                credit = 0.5 * 0.75 * ZONE_MULT[1]
        # stuffing is judged per literal variant (an engine matching "vector search" does not count "embeddings")
        lit_max = max((len(_unique_hits([p], alltext)) for p in pats), default=0)
        flag = lit_max >= 4 and not home
        if flag: stuffed.append(t["term"])
        tot_w += w; tot_credit += w * credit
        rows.append({"term": t["term"], "priority": t.get("priority", "preferred"), "weight": w, "mentions": total,
                     "per_zone": per, "zones_hit": zones_hit, "synonym_hits": syn, "credit": round(credit, 2),
                     "stuffed": flag, "variants": pats, "auto": t.get("auto", False)})
    score = round(tot_credit / tot_w * 30 - min(5, len(stuffed)), 1) if tot_w else None
    return {"terms": rows, "weighted_credit": round(tot_credit, 2), "weight_total": tot_w, "score_of_30": score,
            "stuffing_deduction": min(5, len(stuffed)),
            "missing": [r["term"] for r in rows if r["mentions"] == 0 and not r["synonym_hits"]],
            "synonym_only": [r["term"] for r in rows if r["mentions"] == 0 and r["synonym_hits"]],
            "single_zone": [r["term"] for r in rows if r["zones_hit"] == 1],
            "over_cap": stuffed,
            "at_cap_3": [r["term"] for r in rows if r["mentions"] == 3],
            "no_summary_cap": not bool(zone_text.get("summary", "").strip())}

# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--terms", help="terms.json (see header)")
    ap.add_argument("--jd", help="job posting text/markdown; auto-extracts candidate terms (prune them)")
    ap.add_argument("--out", help="write JSON report here (default: <file>.parse.json next to the input)")
    ap.add_argument("--quiet", action="store_true", help="print only the JSON path, no digest")
    ap.add_argument("--sequential", action="store_true", help="disable parallel extraction (for timing comparisons)")
    a = ap.parse_args()

    t0 = time.time()
    src = os.path.abspath(a.file)
    ext = os.path.splitext(src)[1].lower()
    work = tempfile.mkdtemp(prefix="ats-parse-")
    report = {"file": src, "format": ext.lstrip("."), "compiled_from": None, "warnings": []}
    pdf = None

    if ext == ".tex":
        pdf, err = compile_tex(src, work)
        if err:
            report["warnings"].append(err)
        else:
            report["compiled_from"] = src
    elif ext == ".pdf":
        pdf = src
    elif ext == ".docx":
        d, err = docx_report(src)
        if err:
            report["warnings"].append(err); d = {"stream": "", "layout": ""}
        report.update({k: v for k, v in d.items() if k not in ("stream", "layout")})
        stream, layout = d["stream"], d["layout"]
    elif ext in (".md", ".txt"):
        stream = layout = open(src, encoding="utf-8", errors="ignore").read()
    else:
        print(json.dumps({"error": f"unsupported input {ext}"})); sys.exit(2)

    if pdf:
        for tool in ("pdftotext", "pdffonts", "pdfinfo"):
            if not have(tool):
                report["warnings"].append(f"{tool} not installed (poppler-utils)")
        tasks = {"text": lambda: x_pdftotext(pdf), "layout": lambda: x_layout(pdf), "fonts": lambda: x_fonts(pdf),
                 "pageinfo": lambda: x_pageinfo(pdf), "glyphs": lambda: x_glyphs(pdf),
                 "margins": lambda: x_ink_margins(pdf, work)}
        results = {}
        if a.sequential:
            for k, fn in tasks.items():
                results[k] = fn()
        else:
            with ThreadPoolExecutor(max_workers=len(tasks)) as ex:
                futs = {k: ex.submit(fn) for k, fn in tasks.items()}
                results = {k: f.result() for k, f in futs.items()}
        stream = results["text"].get("stream", ""); layout = results["layout"].get("layout", "")
        fonts = results["fonts"].get("fonts", [])
        report.update({
            "pages": results["pageinfo"].get("pages"), "page_size": results["pageinfo"].get("page_size"),
            "fonts": fonts,
            "type3_fonts": [f["name"] for f in fonts if "Type 3" in f["type"]],
            "fonts_without_unicode_map": [f["name"] for f in fonts if not f["unicode_map"]],
            "fonts_not_embedded": [f["name"] for f in fonts if not f["embedded"]],
            "glyphs": results["glyphs"], "ink_margins_in": results["margins"],
            "text_extracted": bool(stream.strip()),
        })
        for k in ("text", "layout", "fonts", "pageinfo"):
            if results[k].get("error"):
                report["warnings"].append(f"{k}: {results[k]['error']}")

    headings, zones = analyze_headings(stream)
    report.update({
        "contact": analyze_contact(stream),
        "sections": headings,
        "zone_line_counts": {z: len([l for l in ls if l.strip()]) for z, ls in zones.items()},
        "dates": analyze_dates(stream),
        "placeholders": analyze_placeholders(stream),
        "hyphen_joins": analyze_hyphen_joins(stream, layout),
        "column_reorder_suspects": analyze_column_reorder(stream, layout),
        "stream": stream,
        "layout": layout,
    })

    terms = None
    if a.terms:
        terms = json.load(open(a.terms, encoding="utf-8"))
    elif a.jd:
        terms = extract_jd_terms(open(a.jd, encoding="utf-8", errors="ignore").read())
        report["jd_terms_auto_extracted"] = True
    if terms:
        report["keywords"] = count_terms(terms, zones)

    report["elapsed_seconds"] = round(time.time() - t0, 2)
    out = a.out or (src + ".parse.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, ensure_ascii=False)

    if a.quiet:
        print(out); return

    # ---- digest
    r = report
    print(f"parse report → {out}   ({r['elapsed_seconds']}s)")
    if r.get("compiled_from"): print(f"compiled from {os.path.basename(r['compiled_from'])}")
    if pdf:
        g = r["glyphs"]; m = r["ink_margins_in"]
        print(f"pages: {r.get('pages')}   fonts: {len(r['fonts'])} (Type 3: {r['type3_fonts'] or 'none'}; no Unicode map: {r['fonts_without_unicode_map'] or 'none'})")
        print(f"min body size: {g.get('min_body_pt', '?')} pt   families: {g.get('font_families', '?')}")
        if not m.get("error"):
            print(f"ink margins (in): top {m['top_in']} bottom {m['bottom_in']} left {m['left_in']} right {m['right_in']}")
    c = r["contact"]
    print(f"contact: name='{c['name_line']}' email={c['email_in_first_5']} phone={c['phone_in_first_5']} urls={len(c['urls_in_first_5'])}")
    s = r["sections"]
    print(f"sections: {[h['text'] for h in s['headings'] if h['zone']]}  summary={s['has_summary']} skills={s['has_skills']} experience={s['has_experience']} education={s['has_education']}")
    if s["non_canonical_suspects"]: print(f"  non-canonical heading suspects: {s['non_canonical_suspects']}")
    d = r["dates"]
    print(f"dates: {d['styles']}  consistent={d['consistent_month_style']} flags={d['flags'] or 'none'}")
    print(f"placeholders: {r['placeholders'] or 'none'}")
    hj = [j for j in r["hyphen_joins"] if j["in_stream_as"] == "joined"]
    print(f"line-end hyphen joins: {len(r['hyphen_joins'])}; joined without hyphen: {[j['joined'] for j in hj] or 'none'} (real compounds among these are fusions)")
    print(f"column-reorder suspects: {len(r['column_reorder_suspects'])}" + ("" if not r["column_reorder_suspects"] else
          "  e.g. " + json.dumps(r["column_reorder_suspects"][0], ensure_ascii=False)))
    if "keywords" in r:
        k = r["keywords"]
        print(f"keywords: {k['score_of_30']}/30 (stuffing −{k['stuffing_deduction']})  missing={k['missing']}  synonym-only={k['synonym_only']}  single-zone={k['single_zone']}  over-cap={k['over_cap']}  at-3={k['at_cap_3']}" + ("  [no summary → ×0.85 cap]" if k["no_summary_cap"] else ""))
        if r.get("jd_terms_auto_extracted"):
            print("  (terms auto-extracted from the posting — prune false positives and add missing ones via --terms)")
    if r["warnings"]: print("warnings: " + "; ".join(r["warnings"]))

if __name__ == "__main__":
    main()
